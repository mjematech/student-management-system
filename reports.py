import io
from datetime import datetime

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NAVY = colors.HexColor("#1B2A4A")
AMBER = colors.HexColor("#E8A33D")
LIGHT_BG = colors.HexColor("#F2F0EA")


def generate_student_report_pdf(student, marks, average, total_marks,
                                 avg_by_course, attendance_total,
                                 attendance_present, attendance_pct):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        topMargin=2 * cm, bottomMargin=2 * cm,
        leftMargin=2 * cm, rightMargin=2 * cm,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "TitleStyle", parent=styles["Heading1"],
        textColor=NAVY, fontSize=20, spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        "SubtitleStyle", parent=styles["Normal"],
        textColor=colors.HexColor("#6B7280"), fontSize=10, spaceAfter=16,
    )
    section_style = ParagraphStyle(
        "SectionStyle", parent=styles["Heading2"],
        textColor=NAVY, fontSize=13, spaceBefore=16, spaceAfter=8,
    )

    elements = []
    elements.append(Paragraph("Student Report", title_style))
    elements.append(Paragraph(
        f"Generated on {datetime.now().strftime('%d %B %Y')}", subtitle_style
    ))

    info_data = [
        ["Registration No", student["registration_no"]],
        ["Full Name", student["full_name"]],
        ["Programme", student["programme"] or "-"],
        ["Gender", student["gender"]],
        ["Phone", student["phone"] or "-"],
        ["Email", student["email"] or "-"],
    ]
    info_table = Table(info_data, colWidths=[5 * cm, 10 * cm])
    info_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), LIGHT_BG),
        ("TEXTCOLOR", (0, 0), (0, -1), NAVY),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E1D8")),
    ]))
    elements.append(info_table)

    elements.append(Paragraph("Academic Summary", section_style))
    summary_data = [
        ["Overall Average", f"{average}%" if average is not None else "N/A"],
        ["Total Assessments Recorded", str(total_marks)],
        ["Attendance", f"{attendance_present}/{attendance_total} "
                        f"({attendance_pct}%)" if attendance_pct is not None else "No records"],
    ]
    summary_table = Table(summary_data, colWidths=[7 * cm, 8 * cm])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), NAVY),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.white),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E1D8")),
    ]))
    elements.append(summary_table)

    if avg_by_course:
        elements.append(Paragraph("Average Score by Course", section_style))
        course_header = ["Course", "Average Score", "Assessments"]
        course_rows = [[c["course_name"], f'{round(float(c["avg_score"]), 2)}%', str(c["total"])]
                       for c in avg_by_course]
        course_table = Table([course_header] + course_rows, colWidths=[8 * cm, 4 * cm, 3 * cm])
        course_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E1D8")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT_BG]),
        ]))
        elements.append(course_table)

    if marks:
        elements.append(Paragraph("Detailed Marks", section_style))
        marks_header = ["Course", "Semester", "Assessment", "Year", "Score"]
        marks_rows = [[
            m["course_name"], m["semester"], m["assessment_type"],
            str(m["year"]), f'{m["score"]}'
        ] for m in marks]
        marks_table = Table([marks_header] + marks_rows,
                             colWidths=[5 * cm, 3 * cm, 3.5 * cm, 2 * cm, 2 * cm])
        marks_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), AMBER),
            ("TEXTCOLOR", (0, 0), (-1, 0), NAVY),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E1D8")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT_BG]),
        ]))
        elements.append(marks_table)
    else:
        elements.append(Spacer(1, 10))
        elements.append(Paragraph("No marks recorded yet.", styles["Normal"]))

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()


def generate_students_excel(students):
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"

    headers = ["ID", "Registration No", "Full Name", "Gender", "Programme",
               "Phone", "Email", "Address", "Date Registered"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1B2A4A", end_color="1B2A4A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_num, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for s in students:
        ws.append([
            s["id"], s["registration_no"], s["full_name"], s["gender"],
            s["programme"] or "", s["phone"] or "", s["email"] or "",
            s["address"] or "", str(s["date_registered"]),
        ])

    for col_num, header in enumerate(headers, start=1):
        max_len = max(
            [len(str(header))] +
            [len(str(row[col_num - 1])) for row in
             [[s["id"], s["registration_no"], s["full_name"], s["gender"],
               s["programme"] or "", s["phone"] or "", s["email"] or "",
               s["address"] or "", str(s["date_registered"])] for s in students]]
        ) if students else len(str(header))
        ws.column_dimensions[get_column_letter(col_num)].width = max_len + 4

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()